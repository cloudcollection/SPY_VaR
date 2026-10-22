"""Optimize the saved editable deck produced by the external layout builder."""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Pt


def equation(shape, tokens, size):
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.clear()
    for value, baseline in tokens:
        run = paragraph.add_run()
        run.text = value
        run.font.name = "Times New Roman"
        run.font.size = Pt(size if baseline == 0 else size * 0.64)
        if baseline:
            run._r.get_or_add_rPr().set("baseline", str(baseline))


def main(path):
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            properties = getattr(shape._element, "spPr", None)
            if properties is not None:
                for existing in list(properties):
                    if existing.tag in {qn("a:effectLst"), qn("a:effectDag")}:
                        properties.remove(existing)
                properties.append(OxmlElement("a:effectLst"))
            if shape.name.startswith("arrow-"):
                line = shape.line._get_or_add_ln()
                for existing in list(line):
                    if existing.tag == qn("a:tailEnd"):
                        line.remove(existing)
                arrow = OxmlElement("a:tailEnd")
                arrow.set("type", "triangle")
                arrow.set("w", "sm")
                arrow.set("len", "sm")
                line.append(arrow)
            if shape.name.startswith("dashed-"):
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if shape.name == "formula-c2":
                equation(shape, [("q̂", 0), ("C2", 30000), ("α,t", -25000),
                    (" = sort", 0), ("α", -25000), (" [ q̂", 0), ("GARCH", 30000),
                    ("α,t", -25000), (" − s · softplus(g", 0), ("θ", -25000),
                    ("(x))", 0), ("α", -25000), (" ]", 0)], 23)
            if shape.name == "formula-pinball":
                equation(shape, [("L(θ) = mean", 0), ("α,t", -25000), (" ρ", 0),
                    ("α", -25000), ("(r", 0), ("t", -25000), (" − q̂", 0),
                    ("α,t", -25000), (");    ρ", 0), ("α", -25000),
                    ("(u) = max(αu, (α − 1)u)", 0)], 21)
    deck.save(path)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
